from pathlib import Path

from openpyxl import load_workbook

from aa_tool.model import Document
from aa_tool.resources import resource_path

TEMPLATE_NAME = "Template File Documents.xlsx"

COLUMNS = [
    "Source",
    "Assignment Number",
    "Document Group",
    "Bookmark Formula",
    "Index#",
    "Document Type",
    "Received Date",
    "Document Date",
    "Effective Date",
    "Adjudicated Date",
    "Grantor",
    "Grantee",
    "Legal Description",
    "Remarks",
    "Abstract Remarks",
    "Internal Remarks",
    "Action Status",
]


def _as_number(value: str):
    """Return an int when the string is all digits, else the original string."""
    return int(value) if value.isdigit() else value


def export_excel(documents: list[Document], out_path: Path) -> None:
    wb = load_workbook(resource_path(TEMPLATE_NAME))
    ws = wb["Index"]

    for doc in documents:
        r = doc.index + 1  # header occupies row 1
        ws.cell(row=r, column=1, value=_as_number(doc.source))
        ws.cell(row=r, column=2, value=_as_number(doc.assignment))
        ws.cell(
            row=r,
            column=4,
            value=f'=E{r}&"-"&F{r}&"-"&IF(G{r}="","NA",TEXT(G{r},"m/d/yyyy"))',
        )
        ws.cell(row=r, column=5, value=doc.index)
        ws.cell(row=r, column=6, value="Unknown")

    wb.save(out_path)
