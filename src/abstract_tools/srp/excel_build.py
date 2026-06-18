# src/abstract_tools/srp/excel_build.py
"""Write the cleaned Case Actions sheet and a verbatim SRP-sheet copy into the
Abstract Worksheet, saving over the original file.
"""

import os
import tempfile
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from abstract_tools.srp.config import (
    COLUMN_WIDTHS,
    DATE_COLUMNS,
    DATE_FORMAT,
    SHEET_NAME,
)
from abstract_tools.srp.extract import extract_case_actions, read_srp
from abstract_tools.srp.transform import clean_case_actions

_SRP_SHEET = "SRP"


def _write_case_actions(ws, df: pd.DataFrame) -> None:
    """Write the DataFrame plus column widths, wrapping, date format, filter."""
    for col_idx, name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=col_idx, value=name)
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            # NaT/blank dates come through as empty; write "" rather than a marker.
            cell_value = "" if value is pd.NaT or pd.isna(value) else value
            ws.cell(row=row_idx, column=col_idx, value=cell_value)

    for col_idx, name in enumerate(df.columns, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(name, 15)
        for row_idx in range(1, len(df) + 2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_idx > 1 and name in DATE_COLUMNS:
                cell.number_format = DATE_FORMAT

    if len(df) > 0:
        last = get_column_letter(len(df.columns))
        ws.auto_filter.ref = f"A1:{last}{len(df) + 1}"
    ws.freeze_panes = "A2"


def _copy_sheet(source, target) -> None:
    """Copy values + formatting from one worksheet to another."""
    for row in source.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            tgt = target.cell(row=cell.row, column=cell.column)
            tgt.value = cell.value
            if cell.has_style:
                tgt.font = copy(cell.font)
                tgt.border = copy(cell.border)
                tgt.fill = copy(cell.fill)
                tgt.number_format = cell.number_format
                tgt.protection = copy(cell.protection)
                tgt.alignment = copy(cell.alignment)

    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    for letter, dim in source.column_dimensions.items():
        if dim.width:
            target.column_dimensions[letter].width = dim.width
    for idx, dim in source.row_dimensions.items():
        if dim.height:
            target.row_dimensions[idx].height = dim.height


def run_srp_merge(srp_path: Path, worksheet_path: Path) -> None:
    """Add 'SRP Case Actions' and 'SRP' sheets to the worksheet, in place."""
    case_actions = clean_case_actions(extract_case_actions(read_srp(srp_path)))

    workbook = load_workbook(worksheet_path)
    for name in (SHEET_NAME, _SRP_SHEET):
        if name in workbook.sheetnames:
            del workbook[name]

    _write_case_actions(workbook.create_sheet(SHEET_NAME), case_actions)

    srp_source = load_workbook(srp_path, data_only=False).active
    _copy_sheet(srp_source, workbook.create_sheet(_SRP_SHEET))

    fd, tmp_name = tempfile.mkstemp(suffix=".xlsx", dir=worksheet_path.parent)
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        workbook.save(tmp_path)
        os.replace(tmp_path, worksheet_path)
    except BaseException:
        tmp_path.unlink(missing_ok=True)
        raise
