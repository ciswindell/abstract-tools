# tests/test_srp_excel_build.py
from pathlib import Path

from openpyxl import Workbook, load_workbook

from abstract_tools.srp import run_srp_merge
from abstract_tools.srp.config import DATE_FORMAT


def _make_srp(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["Serial Register Page"])
    ws.append(["CASE ACTIONS"])
    ws.append(["Action Date", "Date Filed", "Action Name", "Action Status",
               "Action Information"])
    ws.append(["2020-01-05", "2020-01-01", "Lease Issued", "Active", "Info A"])
    ws.append(["2019-06-10", "2019-06-01", "Application", "Closed", "Info B"])
    wb.save(path)
    return path


def _make_worksheet(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Abstract"
    ws["A1"] = "KEEP ME"
    wb.save(path)
    return path


def test_merge_adds_both_sheets_in_order(tmp_path):
    srp = _make_srp(tmp_path / "srp.xlsx")
    worksheet = _make_worksheet(tmp_path / "abstract.xlsx")

    run_srp_merge(srp, worksheet)

    wb = load_workbook(worksheet)
    assert wb.sheetnames == ["Abstract", "SRP Case Actions", "SRP"]
    # Original content preserved.
    assert wb["Abstract"]["A1"].value == "KEEP ME"
    # Case Actions header row.
    assert wb["SRP Case Actions"]["A1"].value == "Document Type"


def test_case_actions_date_cells_formatted(tmp_path):
    srp = _make_srp(tmp_path / "srp.xlsx")
    worksheet = _make_worksheet(tmp_path / "abstract.xlsx")
    run_srp_merge(srp, worksheet)

    ws = load_workbook(worksheet)["SRP Case Actions"]
    # Column B is "Received Date"; row 2 is the first data row.
    assert ws["B2"].number_format == DATE_FORMAT


def test_rerun_overwrites_without_duplicating_sheets(tmp_path):
    srp = _make_srp(tmp_path / "srp.xlsx")
    worksheet = _make_worksheet(tmp_path / "abstract.xlsx")
    run_srp_merge(srp, worksheet)
    run_srp_merge(srp, worksheet)  # second run must not error or duplicate

    wb = load_workbook(worksheet)
    assert wb.sheetnames == ["Abstract", "SRP Case Actions", "SRP"]


def test_srp_tab_font_raised_to_minimum_11(tmp_path):
    from openpyxl.styles import Font

    # SRP source: a 14pt title (already large), a 9pt note (too small),
    # then the CASE ACTIONS table (default font).
    srp = tmp_path / "srp.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Big Title"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = "Tiny note"
    ws["A2"].font = Font(size=9)
    ws.append(["CASE ACTIONS"])
    ws.append(["Action Date", "Date Filed", "Action Name", "Action Status",
               "Action Information"])
    ws.append(["2020-01-05", "2020-01-01", "Lease Issued", "Active", "Info A"])
    wb.save(srp)

    # Original worksheet: a 9pt cell that must NOT be raised (not the SRP tab).
    worksheet = tmp_path / "abstract.xlsx"
    awb = Workbook()
    aws = awb.active
    aws.title = "Abstract"
    aws["A1"] = "KEEP ME"
    aws["A1"].font = Font(size=9)
    awb.save(worksheet)

    run_srp_merge(srp, worksheet)

    out = load_workbook(worksheet)
    srp_ws = out["SRP"]
    # Larger-than-11 size is left as-is, with other attributes preserved.
    assert srp_ws["A1"].font.size == 14.0
    assert srp_ws["A1"].font.bold is True
    assert srp_ws["A1"].font.name == "Arial"
    # Smaller-than-11 size is raised to the 11 floor.
    assert srp_ws["A2"].font.size == 11.0
    # Default-font cells (the table) stay at 11.
    assert srp_ws["A3"].font.size == 11.0
    # The original worksheet sheet is left untouched (9pt stays 9pt).
    assert out["Abstract"]["A1"].font.size == 9.0


def test_failed_replace_leaves_original_intact(tmp_path, monkeypatch):
    import abstract_tools.srp.excel_build as eb
    srp = _make_srp(tmp_path / "srp.xlsx")
    worksheet = _make_worksheet(tmp_path / "abstract.xlsx")
    original_bytes = worksheet.read_bytes()

    def boom(src, dst):
        raise OSError("simulated replace failure")
    monkeypatch.setattr(eb.os, "replace", boom)

    import pytest
    with pytest.raises(OSError, match="simulated replace failure"):
        run_srp_merge(srp, worksheet)

    # Original worksheet must be untouched, and no temp file left behind.
    assert worksheet.read_bytes() == original_bytes
    xlsx_files = sorted(p.name for p in tmp_path.glob("*.xlsx"))
    assert xlsx_files == ["abstract.xlsx", "srp.xlsx"]
