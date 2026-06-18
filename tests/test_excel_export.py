from pathlib import Path

from openpyxl import load_workbook

from abstract_tools.ingest import SourcePdf
from abstract_tools.model import Document, MergedPage
from abstract_tools.excel_export import export_excel, COLUMNS


def _doc(index, file_number, assignment):
    src = SourcePdf(Path(f"{file_number}.pdf"), file_number, assignment, 1)
    page = MergedPage(index - 1, src, 0)
    return Document(index=index, pages=[page])


def test_export_excel_uses_template_headers_and_fills_rows(tmp_path):
    docs = [_doc(1, "368481", "0"), _doc(2, "368495", "0"), _doc(3, "368495", "0")]
    out = tmp_path / "B11294 File Documents.xlsx"
    export_excel(docs, out)

    wb = load_workbook(out)
    ws = wb["Index"]
    # Headers come from the bundled template, in the exact order.
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.title == "Index"
    # Frozen header preserved from template.
    assert ws.freeze_panes == "A2"
    # Row 2 (first document)
    assert ws.cell(row=2, column=1).value == 368481          # Source (numeric)
    assert ws.cell(row=2, column=2).value == 0               # Assignment Number
    assert ws.cell(row=2, column=4).value == '=E2&"-"&F2&"-"&IF(G2="","NA",TEXT(G2,"m/d/yyyy"))'
    assert ws.cell(row=2, column=5).value == 1               # Index#
    assert ws.cell(row=2, column=6).value == "Unknown"       # Document Type
    assert ws.cell(row=2, column=7).value is None            # Received Date blank
    # Row 4 (third document, same source repeats)
    assert ws.cell(row=4, column=1).value == 368495
    assert ws.cell(row=4, column=5).value == 3
    # No stray data rows beyond the documents.
    assert ws.max_row == 4
