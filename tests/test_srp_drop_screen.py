# tests/test_srp_drop_screen.py
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path

from openpyxl import Workbook, load_workbook

from abstract_tools.ui.srp_parser.drop_screen import DropScreen


def _make_srp(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["Serial Register Page"])
    ws.append(["CASE ACTIONS"])
    ws.append(["Action Date", "Date Filed", "Action Name", "Action Status",
               "Action Information"])
    ws.append(["2020-01-05", "2020-01-01", "Lease Issued", "Active", "Info A"])
    wb.save(path)
    return path


def _make_worksheet(path: Path) -> Path:
    wb = Workbook()
    wb.active.title = "Abstract"
    wb.save(path)
    return path


def test_process_disabled_until_both_files_present(qtbot, tmp_path):
    screen = DropScreen(on_back_to_tools=lambda: None)
    qtbot.addWidget(screen)
    assert not screen.process_button.isEnabled()

    screen.srp_zone.set_path(_make_srp(tmp_path / "srp.xlsx"))
    assert not screen.process_button.isEnabled()

    screen.worksheet_zone.set_path(_make_worksheet(tmp_path / "abs.xlsx"))
    assert screen.process_button.isEnabled()


def test_do_merge_writes_sheets(qtbot, tmp_path):
    srp = _make_srp(tmp_path / "srp.xlsx")
    worksheet = _make_worksheet(tmp_path / "abs.xlsx")
    screen = DropScreen(on_back_to_tools=lambda: None)
    qtbot.addWidget(screen)
    screen.srp_zone.set_path(srp)
    screen.worksheet_zone.set_path(worksheet)

    screen.do_merge()

    wb = load_workbook(worksheet)
    assert "SRP Case Actions" in wb.sheetnames
    assert "SRP" in wb.sheetnames


def test_shutdown_without_run_is_noop(qtbot):
    screen = DropScreen(on_back_to_tools=lambda: None)
    qtbot.addWidget(screen)
    screen.shutdown()  # no worker started — must not raise
